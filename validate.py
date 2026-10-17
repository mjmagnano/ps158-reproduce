"""
Check the numbers behind plot.py, in two parts.

1. FED TABLES. Compare our medians from the microdata with the Fed's own
   published Table 4 ("Family net worth, by selected characteristics of
   families", public data, thousands of 2025 dollars) for every wave.
   The Fed does not publish education x children together, so we check each
   half separately:
     - education of the family head (EDCL 1, 2, 4), all families;
     - families with children (KIDS > 0), split as the Fed does into
       "single with child(ren)" and "couple with child(ren)" (MARRIED).

2. MARGIN OF ERROR for the 2025 medians in the figure, using the Fed's method:
   sampling variance from the 999 replicate weights (implicate 1 only) plus
   (1 + 1/5) x the variance across the 5 implicate estimates. We first check
   that the method reproduces the Fed's published standard errors for
   "All families" and "College degree", then apply it to our subgroups.

Downloads into data/ if missing: the Fed's 2025 historical tables, its
standard-error table, and the 2025 replicate weights (scf2025rw1s.zip).
Needs: pandas, numpy, requests, openpyxl.
"""

import io
import zipfile
from pathlib import Path

import numpy as np
import openpyxl
import pandas as pd
import requests

HERE = Path(__file__).resolve().parent
DATA_DIR = HERE / "data"
FED = "https://www.federalreserve.gov/econres/files/"
WAVES = [1989, 1992, 1995, 1998, 2001, 2004, 2007, 2010, 2013, 2016, 2019, 2022, 2025]
N_REPLICATES = 999

# Our group definition -> the row label in the Fed's Table 4.
TABLE4_CHECKS = {
    "No high school diploma": lambda d: d["EDCL"] == 1,
    "High school diploma": lambda d: d["EDCL"] == 2,
    "College degree": lambda d: d["EDCL"] == 4,
    "Single with child(ren)": lambda d: (d["KIDS"] > 0) & (d["MARRIED"] != 1),
    "Couple with child(ren)": lambda d: (d["KIDS"] > 0) & (d["MARRIED"] == 1),
}


def fetch(filename: str) -> Path:
    path = DATA_DIR / filename
    if not path.exists():
        print(f"downloading {FED + filename}")
        resp = requests.get(FED + filename, timeout=300)
        resp.raise_for_status()
        path.write_bytes(resp.content)
    return path


def replicate_weights_path() -> Path:
    path = DATA_DIR / "p25_rw1.dta"
    if not path.exists():
        with zipfile.ZipFile(fetch("scf2025rw1s.zip")) as zf:
            zf.extract("p25_rw1.dta", DATA_DIR)
    return path


def weighted_median(values: np.ndarray, weights: np.ndarray) -> float:
    keep = weights > 0
    values, weights = values[keep], weights[keep]
    order = np.argsort(values)
    cum = np.cumsum(weights[order])
    return float(values[order][np.searchsorted(cum, cum[-1] / 2)])


def median_of(df: pd.DataFrame) -> float:
    return weighted_median(df["NETWORTH"].to_numpy(), df["WGT"].to_numpy())


def read_table4(filename: str) -> dict:
    """{row label: {year: median in $ thousands}} from the Fed's Table 4."""
    ws = openpyxl.load_workbook(fetch(filename), read_only=True, data_only=True)["Table 4"]
    rows = list(ws.iter_rows(values_only=True))
    header = next(r for r in rows if r[0] == "Family characteristic")
    year_cols = {int(v): i for i, v in enumerate(header) if isinstance(v, int)}  # median column
    return {
        str(r[0]).strip(): {y: r[i] for y, i in year_cols.items()}
        for r in rows if r[0] and isinstance(r[1], (int, float))
    }


def read_published_se() -> dict:
    """(2025 median, 2025 SE) rows from the Fed's standard-error table, in $ thousands.

    In the Fed's file the education rows sit one row above their labels, so
    we return the "All families" row and the whole education block, and
    match on the median instead of the label.
    """
    ws = openpyxl.load_workbook(fetch("2025_tables1_2_stderrs.xlsx"), read_only=True, data_only=True)["Table 2"]
    rows = list(ws.iter_rows(values_only=True))
    start = next(i for i, r in enumerate(rows) if r[0] == "Education of reference person")
    return {
        "All families": [(r[3], r[4]) for r in rows if r[0] == "All families"],
        "College degree": [(r[3], r[4]) for r in rows[start:start + 5] if isinstance(r[3], (int, float))],
    }


def check_fed_tables() -> None:
    print("=" * 78)
    print("1. Our medians vs. the Fed's published Table 4 (public data, $K of 2025)")
    print("=" * 78)
    fed = read_table4("scf2025_tables_public_real_historical.xlsx")
    worst = 0.0
    for year in WAVES:
        df = pd.read_csv(DATA_DIR / f"SCFP{year}.csv", usecols=["EDCL", "NETWORTH", "WGT", "KIDS", "MARRIED"])
        cells = []
        for label, select in TABLE4_CHECKS.items():
            ours = median_of(df[select(df)]) / 1000
            published = fed[label][year]
            worst = max(worst, abs(ours - published))
            cells.append(f"{ours:6.1f}/{published:6.1f}")
        print(f"{year}  " + "  ".join(cells))
    print("columns (ours/Fed): " + ", ".join(TABLE4_CHECKS))
    print(f"largest difference across all {len(WAVES) * len(TABLE4_CHECKS)} cells: ${worst * 1000:,.0f}\n")


def margin_of_error(df: pd.DataFrame, rep: pd.DataFrame) -> tuple:
    """Pooled median, its standard error, plus the replicate and implicate estimates."""
    point = median_of(df)
    implicates = np.array([median_of(g) for _, g in df.groupby("IMPLICATE")])
    first = df[df["IMPLICATE"] == 1]
    weights = rep.loc[first["Y1"].to_numpy()].to_numpy()
    replicates = np.array([
        weighted_median(first["NETWORTH"].to_numpy(), weights[:, r]) for r in range(N_REPLICATES)
    ])
    se = np.sqrt(np.var(replicates) + 1.2 * np.var(implicates, ddof=1))
    return point, se, replicates, implicates


def check_margin_of_error() -> None:
    print("=" * 78)
    print("2. Margin of error, 2025 (Fed method: 999 replicate weights + imputation)")
    print("=" * 78)
    df = pd.read_csv(DATA_DIR / "SCFP2025.csv", usecols=["YY1", "Y1", "EDCL", "NETWORTH", "WGT", "KIDS"])
    df["IMPLICATE"] = df["Y1"] - 10 * df["YY1"]

    rw = pd.read_stata(replicate_weights_path())
    cols = range(1, N_REPLICATES + 1)
    # Replicate weight r = wt1b{r} x mm{r}; missing means not drawn in that replicate.
    rep = pd.DataFrame(
        np.nan_to_num(rw[[f"wt1b{r}" for r in cols]].to_numpy() * rw[[f"mm{r}" for r in cols]].to_numpy()),
        index=rw["y1"].astype(int),
    )

    # Does the method reproduce the Fed's published standard errors? The Fed's
    # SE table uses its internal data, so its medians differ slightly from ours;
    # take the row with the nearest median.
    published = read_published_se()
    print("Method check against the Fed's published SEs ($K):")
    for label, select in (("All families", df["EDCL"] > 0), ("College degree", df["EDCL"] == 4)):
        point, se, _, _ = margin_of_error(df[select], rep)
        fed_median, fed_se = min(published[label], key=lambda p: abs(p[0] - point / 1000))
        print(f"  {label:16s} ours: median {point / 1000:6.1f}, SE {se / 1000:5.1f}"
              f"   Fed: median {fed_median:6.1f}, SE {fed_se:5.1f}")

    print("\nThe figure's 2025 points, families with a child at home:")
    kids = df["KIDS"] > 0
    results = {}
    for label, select in (("College degree", df["EDCL"] == 4), ("High school or less", df["EDCL"] <= 2)):
        sub = df[kids & select]
        point, se, replicates, implicates = margin_of_error(sub, rep)
        results[label] = (point, replicates, implicates)
        print(f"  {label:20s} median ${point:,.0f}  SE ${se:,.0f}  "
              f"95% CI ${point - 1.96 * se:,.0f} to ${point + 1.96 * se:,.0f}  "
              f"({sub['YY1'].nunique()} families in the sample)")

    (c, c_rep, c_imp), (h, h_rep, h_imp) = results.values()
    ratio = c / h
    se = np.sqrt(np.var(c_rep / h_rep) + 1.2 * np.var(c_imp / h_imp, ddof=1))
    print(f"  {'Ratio':20s} {ratio:.1f}x  SE {se:.1f}  95% CI {ratio - 1.96 * se:.1f}x to {ratio + 1.96 * se:.1f}x")


def main() -> int:
    check_fed_tables()
    check_margin_of_error()
    return 0


if __name__ == "__main__":
    raise SystemExit(main())
